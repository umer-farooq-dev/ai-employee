"""CRM-13 lead import: upload CSV / Excel → map columns → validate, normalise phones, detect duplicates → report.

Limits (rows, file size) come from Super Admin → System Settings → CRM. Rows are written in bulk so 10,000-row files
finish within one request; errors are kept per row and offered as a CSV report.
"""

import csv
import io
import os

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone

from apps.audit.services import record
from apps.core.storage import company_path, private_storage
from apps.industries import services as templates
from apps.industries.models import LeadSource
from apps.platform_config import services as config

from . import services as crm
from .models import AssignmentRule, Lead, LeadActivity, LeadImport

CORE_TARGETS = [
    ("name", "Name"),
    ("phone", "Phone"),
    ("email", "Email"),
    ("source", "Source"),
    ("status", "Status"),
    ("tags", "Tags"),
    ("next_action", "Next action"),
    ("follow_up_at", "Follow-up date"),
    ("preferred_language", "Preferred language"),
]
MAX_ERRORS = 5000


def targets(company):
    return CORE_TARGETS + [(f"custom:{d.key}", d.label) for d in templates.fields_for(company, "lead")]


def _read(name, data: bytes):
    ext = os.path.splitext(name)[1].lower()
    if ext == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(text)))
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if c is None else str(c) for c in row] for row in wb.active.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ValidationError("Upload a .csv or .xlsx file.")


def start(company, upload, user=None, request=None) -> LeadImport:
    max_mb = config.get("crm.import_max_mb")
    if upload.size > max_mb * 1024 * 1024:
        raise ValidationError(f"The file is larger than {max_mb} MB.")
    data = upload.read()
    rows = [r for r in _read(upload.name, data) if any(str(c).strip() for c in r)]
    if not rows:
        raise ValidationError("The file is empty.")
    limit = config.get("crm.import_max_rows")
    if len(rows) - 1 > limit:
        raise ValidationError(f"The file has {len(rows) - 1} rows; the limit is {limit}.")
    key = private_storage.save(company_path(company.id, "imports", os.path.basename(upload.name)), ContentFile(data))
    headers = [str(h).strip() for h in rows[0]]
    imp = LeadImport.objects.create(
        company=company,
        file_name=upload.name[:255],
        file_key=key,
        headers=headers,
        total=len(rows) - 1,
        created_by=user,
        mapping=suggest_mapping(company, headers),
    )
    record(request, "lead.import_upload", "lead_import", imp.id, company=company, changes={"rows": imp.total})
    return imp


def suggest_mapping(company, headers):
    by_label = {}
    for key, label in targets(company):
        by_label[label.lower()] = key
        by_label[key.split(":")[-1].lower()] = key
    aliases = {
        "mobile": "phone",
        "mobile number": "phone",
        "phone number": "phone",
        "full name": "name",
        "e-mail": "email",
        "email address": "email",
    }
    return {h: by_label.get(h.lower(), aliases.get(h.lower(), "")) for h in headers}


def _rows(imp):
    with private_storage.open(imp.file_key, "rb") as fh:
        rows = _read(imp.file_name, fh.read())
    return [r for r in rows if any(str(c).strip() for c in r)][1:]


@transaction.atomic
def run(imp, mapping, duplicate_mode="skip", user=None, request=None) -> LeadImport:
    company = imp.company
    valid_targets = {k for k, _ in targets(company)}
    mapping = {h: t for h, t in (mapping or {}).items() if h in imp.headers and t in valid_targets}
    if not any(t in ("name", "phone", "email") for t in mapping.values()):
        raise ValidationError("Map at least one of Name, Phone or Email.")
    index = {h: i for i, h in enumerate(imp.headers)}
    defs = templates.fields_for(company, "lead")
    default_status = crm.default_status(company)
    import_source = LeadSource.objects.filter(key="import").first()
    existing = list(Lead.objects.filter(company=company, deleted_at__isnull=True))
    by_phone = {lead.phone: lead for lead in existing if lead.phone}
    by_email = {lead.email.lower(): lead for lead in existing if lead.email}
    seen_phone, seen_email = set(), set()
    rules = list(AssignmentRule.all_companies.select_for_update().filter(company=company, is_active=True))
    ctx, cache = crm.scoring_context(company), {}
    to_create, errors, updated, skipped = [], [], 0, 0

    for number, row in enumerate(_rows(imp), start=2):  # row 1 is the header line
        raw = {t: (row[index[h]] if index[h] < len(row) else "") for h, t in mapping.items()}
        core_raw = {k: v for k, v in raw.items() if not k.startswith("custom:")}
        custom_raw = {k[7:]: v for k, v in raw.items() if k.startswith("custom:")}
        try:
            core = crm.clean_core(company, core_raw, cache=cache)
            custom = crm.field_engine.clean_values(defs, custom_raw, partial=True)
        except ValidationError as exc:
            errors.append({"row": number, "reason": "; ".join(exc.messages)})
            continue
        phone, email = core.get("phone", ""), (core.get("email") or "").lower()
        if (phone and phone in seen_phone) or (email and email in seen_email):
            skipped += 1
            errors.append({"row": number, "reason": "Duplicate of an earlier row in this file"})
            continue
        seen_phone.add(phone) if phone else None
        seen_email.add(email) if email else None
        match = by_phone.get(phone) if phone else None
        match = match or (by_email.get(email) if email else None)
        if match:
            if duplicate_mode != "update":
                skipped += 1
                errors.append({"row": number, "reason": f"Already exists as {match.reference}"})
                continue
            for key, value in core.items():
                if value not in (None, "", []):
                    setattr(match, key, value)
            match.custom = {**match.custom, **{k: v for k, v in custom.items() if v not in (None, "", [])}}
            match.save()
            LeadActivity.objects.create(
                company=company,
                lead=match,
                type="import",
                summary=f"Updated by import {imp.file_name}"[:500],
                actor=user,
            )
            updated += 1
            continue
        lead = Lead(company=company, custom=custom, created_via="import", **core)
        lead.status = core.get("status") or default_status
        lead.source = core.get("source") or import_source
        lead.score_value, lead.score_label, lead.score_reason, _ = crm.compute_score(lead, company, ctx)
        if not lead.assigned_user_id and rules:
            lead.assigned_user, _ = crm.pick_assignee(company, lead, rules, ctx)
        to_create.append(lead)

    for lead, number in zip(to_create, crm._next_numbers(company, len(to_create)), strict=True):
        lead.number = number
    Lead.objects.bulk_create(to_create, batch_size=1000)
    LeadActivity.objects.bulk_create(
        [
            LeadActivity(
                company=company,
                lead=lead,
                type="import",
                summary=f"Imported from {imp.file_name}"[:500],
                actor=user if user and user.is_authenticated else None,
            )
            for lead in to_create
        ],
        batch_size=1000,
    )
    for rule in rules:
        rule.save(update_fields=["last_index", "updated_at"])

    imp.mapping, imp.duplicate_mode, imp.status = mapping, duplicate_mode, LeadImport.DONE
    imp.created, imp.updated, imp.skipped = len(to_create), updated, skipped
    imp.failed = len([e for e in errors if not e["reason"].startswith(("Duplicate", "Already"))])
    imp.errors = errors[:MAX_ERRORS]
    imp.save()
    record(
        request,
        "lead.import",
        "lead_import",
        imp.id,
        company=company,
        changes={
            "created": imp.created,
            "updated": imp.updated,
            "skipped": imp.skipped,
            "failed": imp.failed,
            "finished_at": timezone.now().isoformat(),
        },
    )
    return imp


def error_report(imp) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Row", "Reason"])
    for e in imp.errors:
        writer.writerow([e["row"], e["reason"]])
    return buf.getvalue().encode("utf-8-sig")
