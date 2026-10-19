"""Read CSV / Excel uploads into rows of strings (lead and catalogue imports)."""

import csv
import io
import os

from django.core.exceptions import ValidationError


def read_table(name, data: bytes):
    """→ list of rows (lists of str); blank rows removed. .csv (UTF-8, with or without BOM) or .xlsx."""
    ext = os.path.splitext(name)[1].lower()
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if c is None else str(c) for c in row] for row in wb.active.iter_rows(values_only=True)]
        wb.close()
    else:
        raise ValidationError("Upload a .csv or .xlsx file.")
    return [r for r in rows if any(str(c).strip() for c in r)]


def write_table(rows, fmt="csv", sheet="Sheet1"):
    """rows → (bytes, content type, extension) as .xlsx or .csv (UTF-8 with BOM so Excel shows Indian scripts)."""
    if fmt == "xlsx":
        from openpyxl import Workbook

        wb = Workbook(write_only=True)
        ws = wb.create_sheet(sheet[:31])
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode("utf-8-sig"), "text/csv", "csv"
